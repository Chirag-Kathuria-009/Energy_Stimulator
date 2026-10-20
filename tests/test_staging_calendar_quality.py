import csv
from datetime import datetime

import duckdb
import openpyxl
import pytest

from ets.config import load_config
from ets.profiles.bdew_staging import write_bdew_long_csv
from ets.profiles.calendar import write_holiday_seed
from ets.quality.checks import run_gold_checks
from ets.transform import dbt_vars

# ---------------------------------------------------------------- BDEW staging


def make_workbook(path, *, hole=False):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name in ("H25", "P25"):
        ws = wb.create_sheet(name)
        for m in range(12):
            for k, daytype in enumerate(("SA", "FT", "WT")):
                col = 3 + 3 * m + k
                ws.cell(3, col, datetime(2012, m + 1, 1))
                ws.cell(4, col, daytype)
                for q in range(96):
                    ws.cell(5 + q, col, 10.0 * (m + 1) + q / 100 + k)
    if hole:
        wb["H25"].cell(20, 10).value = None  # cell(r, c, None) would leave the value unchanged
    wb.save(path)


def test_bdew_staging_writes_one_row_per_profile_month_daytype_quarter_hour(tmp_path):
    xlsx = tmp_path / "bdew.xlsx"
    make_workbook(xlsx)
    out = write_bdew_long_csv(xlsx, tmp_path / "out" / "long.csv")
    rows = list(csv.DictReader(out.open()))
    assert len(rows) == 2 * 36 * 96
    pick = next(
        r
        for r in rows
        if (r["profile_id"], r["month"], r["daytype"], r["qh_of_day"]) == ("H25", "3", "FT", "7")
    )
    assert float(pick["kwh_per_million"]) == pytest.approx(
        10.0 * 3 + 7 / 100 + 1
    )  # column k=1 (FT)


def test_bdew_staging_rejects_a_workbook_with_a_missing_value(tmp_path):
    xlsx = tmp_path / "bdew.xlsx"
    make_workbook(xlsx, hole=True)
    with pytest.raises(ValueError, match="bad value"):
        write_bdew_long_csv(xlsx, tmp_path / "long.csv")


# ---------------------------------------------------------------- holiday seed


def test_holiday_seed_has_national_holidays_only(tmp_path):
    path = write_holiday_seed(2025, 2026, tmp_path / "seeds" / "h.csv")
    rows = list(csv.DictReader(path.open()))
    dates = [r["date"] for r in rows]
    assert dates == sorted(dates) and len(dates) == len(set(dates))
    assert {"2025-10-03", "2025-12-25", "2025-12-26", "2026-01-01", "2026-04-03"} <= set(dates)
    assert "2025-10-31" not in dates  # Reformationstag is state-specific
    assert 9 <= sum(d.startswith("2025") for d in dates) <= 11


# ---------------------------------------------------------------- gold checks


def good_tables(con):
    con.execute(
        """
        create table dim_time as
        select i::integer as idx, ts as ts_utc, ts as ts_local, ts::date as date_local,
               (i % 96)::integer as qh_of_day, 1::integer as weekday, false as is_holiday,
               false as is_dst_day, 1::integer as month, 'winter' as season
        from (select unnest(range(0, 96)) as i, timestamp '2026-01-05 00:00:00' as t0) a,
             lateral (select a.t0 + to_minutes(15 * a.i) as ts)
        """
    )
    con.execute(
        "create table fact_price_qh as select idx, ts_utc, 10.0::double as price_ct_kwh, "
        "100.0::double as price_eur_mwh, 'qh' as source_resolution, false as is_estimated from dim_time"
    )
    con.execute(
        "create table fact_heat_qh as select idx, ts_utc, (1/96.0)::double as heat_space_norm, "
        "(1/96.0)::double as heat_water_norm, 3.0::double as cop_space, 3.0::double as cop_water from dim_time"
    )
    con.execute(
        "create table fact_load_profile_qh as select idx, p as profile_id, 1.0::double as kwh, "
        "false as is_estimated from dim_time, (values ('H25'), ('P25')) t(p)"
    )


@pytest.fixture
def gold():
    con = duckdb.connect()
    good_tables(con)
    return con


def test_valid_gold_tables_pass(gold):
    assert run_gold_checks(gold, n=96) == []


def test_wrong_row_count_is_reported(gold):
    problems = run_gold_checks(gold, n=97)
    assert any("fact_price_qh: 96 rows, expected 97" in p for p in problems)


def test_price_out_of_range_is_reported(gold):
    gold.execute("update fact_price_qh set price_ct_kwh = 250 where idx = 3")
    assert any("price_ct_kwh" in p for p in run_gold_checks(gold, n=96))


def test_cop_out_of_range_is_reported(gold):
    gold.execute("update fact_heat_qh set cop_space = 12 where idx = 0")
    assert any("cop_space" in p for p in run_gold_checks(gold, n=96))


def test_null_price_is_reported(gold):
    gold.execute("update fact_price_qh set price_ct_kwh = null where idx = 5")
    assert any("price_ct_kwh" in p for p in run_gold_checks(gold, n=96))


def test_missing_table_is_reported():
    con = duckdb.connect()
    good_tables(con)
    con.execute("drop table fact_heat_qh")
    assert any("fact_heat_qh: cannot read table" in p for p in run_gold_checks(con, n=96))


def test_profile_with_missing_rows_is_reported(gold):
    gold.execute("delete from fact_load_profile_qh where profile_id = 'P25' and idx = 7")
    assert any("fact_load_profile_qh[P25]: 95 rows" in p for p in run_gold_checks(gold, n=96))


# ---------------------------------------------------------------- dbt vars


def test_dbt_vars_come_from_config_and_settings():
    v = dbt_vars(load_config())
    assert v["sim_start_utc"] == "2025-09-30 22:00:00" and v["sim_end_utc"] == "2026-09-30 22:00:00"
    assert v["n_qh"] == 35040
    assert v["opsd_channels"]["pv_r1"] == "DE_KN_residential1_pv"
    assert set(v["when2heat_columns"]) == {
        "heat_profile_space", "heat_profile_water", "cop_space", "cop_water",
    }  # fmt: skip
    assert v["when2heat_season_start_year"] == 2014
    assert "\\" not in v["bronze_dir"]  # forward slashes, safe inside SQL on Windows
